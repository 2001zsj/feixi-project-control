from pathlib import Path
import json, zipfile, xml.etree.ElementTree as ET
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
out=root/'audit-evidence'
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
results=[]
for name,count in [('selection',16),('difficult',11),('pending',8),('building',13),('complete',27)]:
    file=out/(name+'.xlsx')
    with zipfile.ZipFile(file) as z:
        assert z.testzip() is None
        for entry in z.namelist():
            if entry.endswith('.xml') or entry.endswith('.rels'): ET.fromstring(z.read(entry))
        sheet=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        setup=sheet.find('m:pageSetup',ns)
        assert setup.attrib['paperSize']=='9' and setup.attrib['orientation']=='landscape'
        assert setup.attrib['fitToWidth']=='1' and setup.attrib['fitToHeight']=='0'
    wb=load_workbook(file,data_only=False)
    ws=wb.active
    assert ws.max_row==count+3 and ws.freeze_panes=='C4'
    assert ws.print_title_rows=='$1:$3' and ws.print_area
    headers=[c.value for c in ws[3]]
    for row in ws.iter_rows(min_row=4):
        for c in row:
            assert c.border.left.style=='thin' and c.border.right.style=='thin'
            assert c.font.name=='Microsoft YaHei'
            assert c.data_type!='f', 'User values must not become formulas'
        assert row[1].alignment.wrap_text and row[2].alignment.wrap_text
        assert row[3].data_type=='s', 'Order number must remain text'
    if name=='selection':
        assert headers[-3:]==['最新进展','最近更新时间','跟进状态']
        assert '当前问题' not in headers
        assert ws.cell(4,3).value=='暂无'
    if name=='difficult': assert headers[-4:]==['当前问题','问题状态','最近更新时间','跟进状态']
    if name=='pending':
        assert '选址立项情况' in headers
        values=[dict(zip(headers,r)) for r in ws.iter_rows(min_row=4,values_only=True)]
        target=next(r for r in values if r['需求站名']=='玉兰大道与黄岗路交口-NR3.5G')
        assert target['当前状态']=='待电信确认'
        assert '待电信确认后立项' in target['选址立项情况']
    if name=='building':
        for h in ['施工问题','问题说明','节点提醒','48工作日截止','当前工作日','剩余工作日','工期状态']:assert h in headers
        values=[dict(zip(headers,r)) for r in ws.iter_rows(min_row=4,values_only=True)]
        target=next(r for r in values if r['铁塔站名']=='肥西紫蓬镇宏德苗圃')
        assert target['48工作日截止'].strftime('%Y-%m-%d')=='2026-10-30'
        node_col=headers.index('节点提醒')+1
        for r in range(4,ws.max_row+1):
            if ws.cell(r,node_col).value!='无':assert ws.cell(r,node_col).fill.fgColor.rgb=='FFFFF2CC'
    if name=='complete':assert headers[-1]=='当前状态' and '立项日期' not in headers
    results.append({'module':name,'rows':count,'columns':len(headers),'freeze':ws.freeze_panes,'printArea':str(ws.print_area),'pass':True})
    wb.close()
(out/'export-results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(results,ensure_ascii=False,indent=2))
